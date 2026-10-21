"""User-selected exports with complete JD text, immutable generations and safe cells."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any
from urllib.parse import urlparse, urlunparse
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from job_radar.config import FIELD_LABELS, TaskConfig
from job_radar.store import JobStore, now

DETAIL_LABELS = {"not_requested": "未选择采集", "pending": "待采集", "complete": "已采集", "fetched": "已采集",
                 "empty": "详情未返回正文", "failed": "未能采集", "budget_skipped": "额度不足，未采集", "skipped": "未通过初筛"}
CHANGE_LABELS = {"first_seen": "首次收录", "changed": "信息变化", "seen": "历史已见"}


def safe_job_url(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    try:
        parts = urlparse(value)
    except ValueError:
        return ""
    if parts.scheme != "https" or parts.netloc not in {"www.zhipin.com", "zhipin.com"}:
        return ""
    if not parts.path.startswith("/job_detail/"):
        return ""
    return urlunparse((parts.scheme, parts.netloc, parts.path, "", "", ""))


def public_job(job: dict) -> dict:
    allowed = set(FIELD_LABELS) | {"job_key", "notes"}
    data = {key: value for key, value in job.items() if key in allowed}
    data["url"] = safe_job_url(data.get("url"))
    return data


def _display(field: str, value: Any) -> str:
    if field == "detail_status":
        return DETAIL_LABELS.get(str(value), str(value or ""))
    if field == "change":
        return CHANGE_LABELS.get(str(value), str(value or ""))
    if field == "matched":
        return "保留" if value else "已筛除"
    if isinstance(value, list):
        return "、".join(str(item) for item in value)
    return str(value if value is not None else "")


def _csv_safe(value: str) -> str:
    # Spreadsheets may execute formula-like values even when CSV-quoted.
    if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r")):
        return "'" + value
    return value


def export_run(store: JobStore, run_id: str, config: TaskConfig | None = None) -> list[dict]:
    run = store.get_run(run_id)
    if run is None:
        raise KeyError(run_id)
    config = config or TaskConfig.from_dict(run["config"])
    jobs = store.get_jobs(run_id)
    scope = config.export["scope"]
    if scope == "matched":
        jobs = [job for job in jobs if job.get("matched", True)]
    elif scope == "new":
        jobs = [job for job in jobs if job.get("change") == "first_seen"]
    fields = config.export["fields"]
    selected = [{field: public_job(job).get(field, "") for field in fields} for job in jobs]
    generation = uuid.uuid4().hex[:12]
    output = store.root / "exports" / run_id / generation
    output.mkdir(parents=True, exist_ok=False)
    prefix = "示例岗位" if run.get("demo") else "岗位清单"
    metadata = {"name": run["config"]["name"], "run_id": run_id, "status": run["status"], "demo": bool(run.get("demo")),
                "generated_at": now(), "scope": scope, "rows": len(jobs), "counts": run.get("counts", {}),
                "search": {key: run["config"][key] for key in ("keywords", "cities", "filters", "details")},
                "fields": fields,
                "message": run.get("message", ""), "note": "首次收录为本地首次观察时间，不等于平台发布日期。"}
    paths: list[Path] = []
    for fmt in config.export["formats"]:
        path = output / f"{prefix}.{fmt}"
        if fmt == "json":
            path.write_text(json.dumps({"metadata": metadata, "jobs": selected}, ensure_ascii=False, indent=2), encoding="utf-8")
        elif fmt == "csv":
            with path.open("w", newline="", encoding="utf-8-sig") as stream:
                writer = csv.writer(stream)
                writer.writerow([FIELD_LABELS[field] for field in fields])
                writer.writerows([_csv_safe(_display(field, row[field])) for field in fields] for row in selected)
        else:
            book = Workbook()
            sheet = book.active
            sheet.title = "岗位清单"
            sheet.append([FIELD_LABELS[field] for field in fields])
            long_text: list[tuple[int, str, str]] = []
            for row_index, row in enumerate(selected, 2):
                for col, field in enumerate(fields, 1):
                    text = _display(field, row[field])
                    if len(text) > 32767:
                        long_text.append((row_index, FIELD_LABELS[field], text))
                        text = text[:32000] + "\n（超长原文分段保存在「长文本」工作表）"
                    cell = sheet.cell(row_index, col, text)
                    cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=field in {"description", "notes", "filter_reason"})
                    if field == "url" and safe_job_url(text):
                        cell.hyperlink = text
                        cell.font = Font(color="147D72", underline="single")
                sheet.row_dimensions[row_index].height = 42
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for i, field in enumerate(fields, 1):
                cell = sheet.cell(1, i)
                cell.fill = PatternFill("solid", fgColor="163B38")
                cell.font = Font(color="FFFFFF", bold=True)
                sheet.column_dimensions[get_column_letter(i)].width = 60 if field == "description" else (32 if field in {"title", "company", "filter_reason", "url"} else 20)
            if long_text:
                extra = book.create_sheet("长文本")
                extra.append(["岗位表行号", "字段", "分段", "原文"])
                for row_index, field, text in long_text:
                    for part, offset in enumerate(range(0, len(text), 32000), 1):
                        extra.append([row_index, field, part, text[offset:offset+32000]])
                        extra.cell(extra.max_row, 4).data_type = "s"
            meta = book.create_sheet("运行说明")
            for key, value in metadata.items():
                meta.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)])
                meta.cell(meta.max_row, 2).data_type = "s"
            meta.column_dimensions["A"].width = 24
            meta.column_dimensions["B"].width = 85
            book.save(path)
        paths.append(path)
    # A companion manifest keeps partial/complete status visible even for plain CSV exports.
    manifest = output / "运行说明.json"
    manifest.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    paths.append(manifest)
    exports = [{"name": path.name, "relative": path.relative_to(store.root / "exports").as_posix()} for path in paths]
    store.update_run(run_id, exports=exports)
    return exports
