from __future__ import annotations

import hashlib
import json
import sqlite3
from collections import Counter
from contextlib import closing
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook

from boss_agent_cli.radar.profile import RadarProfile
from boss_agent_cli.radar.evidence_scope import resolve_observation_run_ids


_CHINA_TZ = timezone(timedelta(hours=8))
_EXPECTED_SHEETS = ["今日新增", "本次扫描", "扫描摘要"]
_EXPECTED_JOB_HEADERS = [
	"状态",
	"首次发现时间",
	"岗位",
	"公司",
	"城市/区域",
	"薪资",
	"经验",
	"学历",
	"AI相关证据",
	"BOSS链接",
]


class RadarAcceptanceError(RuntimeError):
	"""Raised when a radar acceptance report cannot be produced."""


def verify_latest_radar_run(
	*,
	database_path: Path,
	workbook_path: Path,
	profile: RadarProfile,
	expected_mode: str = "any",
	require_natural_end: bool = False,
	output_scope: str | None = None,
) -> dict[str, Any]:
	"""Verify the latest radar database run against its exported workbook without writing either input."""
	if expected_mode not in {"any", "baseline", "daily"}:
		raise ValueError(f"unsupported expected mode: {expected_mode}")
	if output_scope is not None and (
		not isinstance(output_scope, str)
		or output_scope not in {"configured_profile", "temporary_page_override"}
	):
		raise ValueError(f"unsupported output scope: {output_scope}")
	database_path = database_path.expanduser().resolve()
	workbook_path = workbook_path.expanduser().resolve()
	if not database_path.is_file():
		raise RadarAcceptanceError(f"岗位雷达数据库不存在: {database_path}")
	if not workbook_path.is_file():
		raise RadarAcceptanceError(f"岗位雷达 Excel 不存在: {workbook_path}")

	database_hash_before = _sha256(database_path)
	workbook_hash_before = _sha256(workbook_path)
	checks: list[dict[str, Any]] = []
	expected_searches = {
		f"{city.name} / {query}"
		for city in profile.cities
		for query in profile.queries
	}

	def check(name: str, passed: bool, details: Any) -> None:
		checks.append({"name": name, "passed": bool(passed), "details": details})

	with closing(_read_only_connection(database_path)) as connection:
		latest, stats, stats_valid = _select_run_for_verification(
			connection,
			profile=profile,
			output_scope=output_scope,
		)
		run_id = str(latest["run_id"])
		mode = str(latest["mode"])

		event_counts = _event_counts(connection, run_id)
		completed_at = _parse_run_time(latest["completed_at"] or latest["started_at"])
		day_start_local = completed_at.astimezone(_CHINA_TZ).replace(hour=0, minute=0, second=0, microsecond=0)
		day_start = day_start_local.astimezone(timezone.utc)
		today_event_counts = _today_event_counts(
			connection,
			profile_fingerprint=profile.fingerprint,
			start=day_start,
			end=completed_at,
		)

		duplicate_identity_groups = _scalar_count(
			connection,
			"SELECT COUNT(*) FROM ("
			"SELECT platform, identity_key FROM radar_jobs WHERE identity_key <> '' "
			"GROUP BY platform, identity_key HAVING COUNT(*) > 1)",
		)
		duplicate_platform_groups = _scalar_count(
			connection,
			"SELECT COUNT(*) FROM ("
			"SELECT platform, platform_job_id FROM radar_jobs WHERE platform_job_id <> '' "
			"GROUP BY platform, platform_job_id HAVING COUNT(*) > 1)",
		)
		security_identifiers = _scalar_count(
			connection,
			"SELECT COUNT(*) FROM radar_jobs WHERE COALESCE(security_id, '') <> '' "
			"OR payload_json LIKE '%securityId%'",
		)
		running_runs = _scalar_count(connection, "SELECT COUNT(*) FROM radar_runs WHERE status = 'running'")
		new_event_first_run_mismatches = _scalar_count(
			connection,
			"SELECT COUNT(*) FROM radar_events AS e JOIN radar_jobs AS j ON j.id = e.job_pk "
			"WHERE e.run_id = ? AND e.event_type IN ('new', 'reposted') AND j.first_run_id <> ?",
			(run_id, run_id),
		)
		configured_job_count = _configured_job_count(connection, profile)
		observation_table_present = _table_exists(connection, "radar_run_observations")
		raw_observation_run_ids = stats.get("observation_run_ids")
		has_logical_observation_scope = isinstance(raw_observation_run_ids, list) and bool(raw_observation_run_ids)
		observation_run_ids = resolve_observation_run_ids(connection, latest, stats)
		logical_event_counts = (
			_event_counts_for_runs(connection, observation_run_ids)
			if has_logical_observation_scope
			else event_counts
		)
		raw_event_window = stats.get("event_window")
		rolling_window_consistent = raw_event_window is None
		rolling_window_details: dict[str, Any] = {"actual": raw_event_window}
		if isinstance(raw_event_window, dict):
			try:
				window_run_ids = raw_event_window.get("run_ids")
				if not isinstance(window_run_ids, list) or tuple(str(value) for value in window_run_ids) != observation_run_ids:
					raise ValueError("run_ids 与逻辑扫描范围不符")
				marks = ",".join("?" for _ in observation_run_ids)
				rows = connection.execute(
					f"SELECT run_id, started_at, completed_at FROM radar_runs WHERE run_id IN ({marks})", observation_run_ids
				).fetchall()
				if len(rows) != len(observation_run_ids):
					raise ValueError("逻辑扫描分段不完整")
				starts = [_parse_run_time(row[1]) for row in rows]
				ends = [_parse_run_time(row[2] or row[1]) for row in rows]
				expected_start, expected_end = min(starts), max(ends)
				rolling_window_consistent = (
					raw_event_window.get("mode") == "rolling"
					and _parse_run_time(raw_event_window.get("start_at")) == expected_start
					and _parse_run_time(raw_event_window.get("end_at")) == expected_end
					and bool(raw_event_window.get("crosses_midnight"))
					== (expected_start.astimezone(_CHINA_TZ).date() != expected_end.astimezone(_CHINA_TZ).date())
				)
				rolling_window_details.update({"expected_run_ids": observation_run_ids, "expected_start_at": expected_start.isoformat(), "expected_end_at": expected_end.isoformat()})
			except (TypeError, ValueError) as exc:
				rolling_window_consistent = False
				rolling_window_details["error"] = str(exc)
		if observation_table_present:
			observation_marks = ",".join("?" for _ in observation_run_ids)
			latest_observed_jobs = _scalar_count(
				connection,
				f"SELECT COUNT(DISTINCT job_pk) FROM radar_run_observations "
				f"WHERE run_id IN ({observation_marks})",
				observation_run_ids,
			)
			first_seen_by_job_id = {
				str(row[0]): _format_first_seen(str(row[1]))
				for row in connection.execute(
					f"SELECT DISTINCT j.platform_job_id, j.first_seen_at FROM radar_jobs AS j "
					f"JOIN radar_run_observations AS o ON o.job_pk = j.id "
					f"WHERE o.run_id IN ({observation_marks}) AND j.platform_job_id <> ''",
					observation_run_ids,
				)
			}
			run_observation_jobs = _scalar_count(
				connection,
				f"SELECT COUNT(DISTINCT job_pk) FROM radar_run_observations "
				f"WHERE run_id IN ({observation_marks})",
				observation_run_ids,
			)
			run_observation_rows = _scalar_count(
				connection,
				f"SELECT COUNT(*) FROM radar_run_observations WHERE run_id IN ({observation_marks})",
				observation_run_ids,
			)
			unpositioned_observations = _scalar_count(
				connection,
				"SELECT COUNT(*) FROM radar_run_observations "
				f"WHERE run_id IN ({observation_marks}) AND (page <= 0 OR rank <= 0)",
				observation_run_ids,
			)
		else:
			latest_observed_jobs = _scalar_count(
				connection,
				"SELECT COUNT(*) FROM radar_jobs WHERE last_run_id = ?",
				(run_id,),
			)
			first_seen_by_job_id = {
				str(row[0]): _format_first_seen(str(row[1]))
				for row in connection.execute(
					"SELECT platform_job_id, first_seen_at FROM radar_jobs "
					"WHERE last_run_id = ? AND platform_job_id <> ''",
					(run_id,),
				)
			}
			run_observation_jobs = 0
			run_observation_rows = 0
			unpositioned_observations = 0

		page_evidence = _natural_end_page_evidence(
			connection,
			run_ids=observation_run_ids,
			expected_searches=expected_searches,
			require_visible_state=stats.get("boss_search_transport") == "page_capture",
		)

	completed_searches = stats.get("completed_searches", [])
	failed_searches = stats.get("failed_searches", [])
	truncated_searches = stats.get("truncated_searches", [])
	completed_search_set = set(completed_searches) if isinstance(completed_searches, list) else set()
	coverage_details = {
		"planned": len(expected_searches),
		"completed": len(completed_search_set),
		"missing": sorted(expected_searches - completed_search_set),
		"unexpected": sorted(completed_search_set - expected_searches),
		"failed": failed_searches,
		"truncated": truncated_searches,
	}
	# Window/file consistency is not evidence that all results were exhausted.
	# A natural-end claim must also retain raw, per-page platform pagination facts.
	raw_page_caps = stats.get("page_cap_reached_searches")
	valid_page_caps = isinstance(raw_page_caps, list) and all(isinstance(item, str) for item in raw_page_caps)
	raw_terminal_evidence_missing = stats.get("terminal_evidence_missing_searches")
	valid_terminal_evidence_metadata = (
		isinstance(raw_terminal_evidence_missing, list)
		and all(isinstance(item, str) for item in raw_terminal_evidence_missing)
	)
	page_capture_terminal_evidence_verified = (
		stats.get("boss_search_transport") != "page_capture"
		or (
			valid_terminal_evidence_metadata
			and not raw_terminal_evidence_missing
		)
	)
	visible_list_end_verified = (
		latest["status"] == "completed"
		and stats.get("complete") is True
		and completed_search_set == expected_searches
		and isinstance(failed_searches, list) and not failed_searches
		and isinstance(truncated_searches, list) and not truncated_searches
		and valid_page_caps and not raw_page_caps
		and page_capture_terminal_evidence_verified
		and not stats.get("stop_code") and not stats.get("stop_reason")
		and page_evidence["proven"]
	)
	if require_natural_end:
		check("visible_list_end_verified", visible_list_end_verified, {
			"page_cap_metadata_valid": valid_page_caps,
			"page_cap_reached_searches": raw_page_caps,
			"terminal_evidence_metadata_valid": valid_terminal_evidence_metadata,
			"terminal_evidence_missing_searches": raw_terminal_evidence_missing,
			"coverage": coverage_details,
			"stop_code": stats.get("stop_code"),
			"stop_reason": stats.get("stop_reason"),
			"page_evidence": page_evidence,
		})

	check("expected_mode", expected_mode == "any" or mode == expected_mode, {"expected": expected_mode, "actual": mode})
	check("stats_json_valid", stats_valid, latest["stats_json"])
	check("rolling_event_window_consistent", rolling_window_consistent, rolling_window_details)
	check(
		"run_completed",
		latest["status"] == "completed" and bool(stats.get("complete")),
		{"status": latest["status"], "complete": stats.get("complete")},
	)
	check("profile_fingerprint", stats.get("profile_fingerprint") == profile.fingerprint, stats.get("profile_fingerprint"))
	if profile.query_page_limits:
		check(
			"configured_search_page_limits",
			stats.get("search_page_limits") == profile.search_page_limits,
			{"expected": profile.search_page_limits, "actual": stats.get("search_page_limits")},
		)
	check(
		"configured_search_coverage",
		completed_search_set == expected_searches and not failed_searches and not truncated_searches,
		coverage_details,
	)
	if stats.get("boss_search_transport") == "page_capture":
		check(
			"page_capture_terminal_evidence",
			page_capture_terminal_evidence_verified,
			{
				"metadata_valid": valid_terminal_evidence_metadata,
				"missing_searches": raw_terminal_evidence_missing,
			},
		)
	check("no_running_runs", running_runs == 0, running_runs)
	check("no_duplicate_identity", duplicate_identity_groups == 0, duplicate_identity_groups)
	check("no_duplicate_platform_job_id", duplicate_platform_groups == 0, duplicate_platform_groups)
	check("no_security_identifiers", security_identifiers == 0, security_identifiers)
	check("new_events_are_first_seen_this_run", new_event_first_run_mismatches == 0, new_event_first_run_mismatches)
	scan_totals = stats.get("scan_totals")
	if isinstance(scan_totals, dict) and has_logical_observation_scope:
		stored_event_counts = scan_totals.get("event_counts")
		check(
			"logical_scan_totals_consistent",
			_as_int(scan_totals.get("unique_seen")) == latest_observed_jobs
			and isinstance(stored_event_counts, dict)
			and {str(key): _as_int(value) for key, value in stored_event_counts.items()} == logical_event_counts,
			{
				"stored_unique_seen": scan_totals.get("unique_seen"),
				"database_unique_seen": latest_observed_jobs,
				"stored_event_counts": stored_event_counts,
				"database_event_counts": logical_event_counts,
			},
		)

	workbook_data = _inspect_workbook(workbook_path, first_seen_by_job_id)
	summary = workbook_data["summary"]
	today_links = workbook_data["today_links"]
	observed_links = workbook_data["observed_links"]
	rolling_window = stats.get("event_window") if isinstance(stats.get("event_window"), dict) else None
	expected_sheets = ["本轮首次发现", "本次扫描", "扫描摘要"] if rolling_window else _EXPECTED_SHEETS
	reported_event_counts = logical_event_counts if rolling_window else today_event_counts
	reported_new_label = "本轮累计新增岗位" if rolling_window else "今日累计新增岗位"
	reported_repost_label = "本轮累计疑似重发" if rolling_window else "今日累计疑似重发"
	# New scans explicitly scope the workbook to the logical scan's run segments,
	# excluding stale historical jobs. Immutable workbooks from older versions did
	# not record that scope and are verified against their historical inventory.
	expected_observed_jobs = (
		configured_job_count
		if mode == "baseline" and bool(stats.get("complete")) and not has_logical_observation_scope
		else latest_observed_jobs
	)
	check("workbook_sheets", workbook_data["sheet_names"] == expected_sheets, workbook_data["sheet_names"])
	check("job_sheet_headers", not workbook_data["header_mismatches"], workbook_data["header_mismatches"])
	check("summary_run_id", summary.get("运行ID") == run_id, summary.get("运行ID"))
	expected_mode_label = (
		"配置变更基线"
		if mode == "baseline" and stats.get("baseline_reason") == "profile_changed"
		else ("首次基线" if mode == "baseline" else "每日增量")
	)
	check("summary_mode", summary.get("运行模式") == expected_mode_label, summary.get("运行模式"))
	page_cap_searches = stats.get("page_cap_reached_searches", [])
	if not isinstance(page_cap_searches, list):
		page_cap_searches = []
	terminal_full_searches = stats.get("terminal_full_searches", [])
	if not isinstance(terminal_full_searches, list):
		terminal_full_searches = []
	terminal_evidence_missing_searches = stats.get("terminal_evidence_missing_searches", [])
	if not isinstance(terminal_evidence_missing_searches, list):
		terminal_evidence_missing_searches = []
	if not stats.get("complete"):
		expected_coverage_label = "配置搜索未完成"
	elif page_cap_searches:
		expected_coverage_label = "配置搜索已完成（部分到安全上限，页面仍有更多）"
	elif terminal_full_searches:
		expected_coverage_label = "配置搜索已完成（终页满载，覆盖待核）"
	elif terminal_evidence_missing_searches:
		expected_coverage_label = "配置搜索已完成（页面终点证据不足）"
	else:
		expected_coverage_label = "配置搜索已完成（页面到当前列表终点；未证明全城搜全）"
	# Older workbooks used the generic "完整" label while still exposing the
	# page-cap count and conclusion. Keep those immutable archives verifiable;
	# newly generated workbooks use the explicit distinction above.
	accepted_coverage_labels = {expected_coverage_label}
	if page_cap_searches:
		accepted_coverage_labels.update({"完整", "观察窗口完整（平台仍有更多）", "已扫描至安全上限（平台仍有更多）"})
	elif terminal_full_searches:
		accepted_coverage_labels.add("完成，但终页满载待验证")
	elif terminal_evidence_missing_searches:
		accepted_coverage_labels.add("完成，但终页证据不足")
	else:
		accepted_coverage_labels.add("完整")
	check(
		"summary_coverage",
		summary.get("覆盖状态") in accepted_coverage_labels,
		{"expected": expected_coverage_label, "actual": summary.get("覆盖状态")},
	)
	check(
		"summary_page_cap_count",
		_as_int(summary.get("达到分页上限的搜索")) == len(page_cap_searches),
		{"expected": len(page_cap_searches), "actual": summary.get("达到分页上限的搜索")},
	)
	check("summary_city_count", _as_int(summary.get("配置城市")) == len(profile.cities), summary.get("配置城市"))
	check(
		"summary_run_new_count",
		_as_int(summary.get("本次新增岗位")) == logical_event_counts["new"],
		summary.get("本次新增岗位"),
	)
	check(
		"summary_run_repost_count",
		_as_int(summary.get("本次疑似重发")) == logical_event_counts["reposted"],
		summary.get("本次疑似重发"),
	)
	check(
		"summary_reported_new_count",
		_as_int(summary.get(reported_new_label)) == reported_event_counts["new"],
		summary.get(reported_new_label),
	)
	check(
		"summary_reported_repost_count",
		_as_int(summary.get(reported_repost_label)) == reported_event_counts["reposted"],
		summary.get(reported_repost_label),
	)
	# Immutable baseline workbooks created before the observation-scope marker
	# listed the baseline inventory on "今日新增". New exports leave that sheet
	# informational-only and put the inventory on "本次扫描".
	expected_today_links = (
		_as_int(summary.get("当前清单岗位"))
		if mode == "baseline" and not has_logical_observation_scope
		else reported_event_counts["new"] + reported_event_counts["reposted"]
	)
	check("today_link_count", len(today_links) == expected_today_links, {"expected": expected_today_links, "actual": len(today_links)})
	check("today_links_unique", len(today_links) == len(set(today_links)), len(today_links))
	check("today_links_valid", all(_is_job_link(target) for target in today_links), len(today_links))
	check(
		"observed_link_count",
		len(observed_links)
		== _as_int(summary.get("当前清单岗位"))
		== expected_observed_jobs,
		{
			"workbook": len(observed_links),
			"summary": summary.get("当前清单岗位"),
			"database": expected_observed_jobs,
		},
	)
	check("observed_links_unique", len(observed_links) == len(set(observed_links)), len(observed_links))
	check("observed_links_valid", all(_is_job_link(target) for target in observed_links), len(observed_links))
	check("no_excel_formulas", not workbook_data["formula_cells"], workbook_data["formula_cells"])
	if mode == "daily":
		check(
			"source_page_rank_recorded",
			observation_table_present
			and run_observation_jobs == latest_observed_jobs
			and run_observation_jobs == len(observed_links)
			and run_observation_rows >= latest_observed_jobs
			and unpositioned_observations == 0,
			{
				"table_present": observation_table_present,
				"observation_jobs": run_observation_jobs,
				"observation_rows": run_observation_rows,
				"unpositioned_rows": unpositioned_observations,
				"latest_observed_jobs": latest_observed_jobs,
				"workbook_rows": len(observed_links),
			},
		)
		check(
			"observed_first_seen_matches_database",
			not workbook_data["first_seen_mismatches"] and len(observed_links) == len(first_seen_by_job_id),
			{
				"mismatches": workbook_data["first_seen_mismatches"][:10],
				"workbook_rows": len(observed_links),
				"database_rows": len(first_seen_by_job_id),
			},
		)

	database_hash_after = _sha256(database_path)
	workbook_hash_after = _sha256(workbook_path)
	inputs_unchanged = database_hash_before == database_hash_after and workbook_hash_before == workbook_hash_after
	check(
		"inputs_unchanged",
		inputs_unchanged,
		{
			"database": {"before": database_hash_before, "after": database_hash_after},
			"workbook": {"before": workbook_hash_before, "after": workbook_hash_after},
		},
	)
	failed_checks = [item["name"] for item in checks if not item["passed"]]
	return {
		"passed": not failed_checks,
		"verification_scope": "natural_end" if require_natural_end else "configured_window",
		# Retain the old key for consumers, but a browser result-list endpoint is
		# never a proof of the platform's complete inventory.
		"platform_results_exhausted": False,
		"city_inventory_exhausted": False,
		"city_inventory_exhaustion_status": "unproven",
		"visible_list_end_verified": visible_list_end_verified,
		"scope_note": "仅核验已采集配置窗口、当前列表终点及文件一致性；不证明全城搜全或平台全量召回。",
		"legacy_coverage_wording": summary.get("覆盖状态") != expected_coverage_label,
		"natural_end_evidence": page_evidence,
		"profile": profile.name,
		"run_selection_scope": (
			"global_latest" if output_scope is None else "profile_fingerprint_and_output_scope"
		),
		"output_scope": stats.get("output_scope", "configured_profile"),
		"mode": mode,
		"run_id": run_id,
		"run_status": str(latest["status"]),
		"database": str(database_path),
		"workbook": str(workbook_path),
		"database_sha256": database_hash_after,
		"workbook_sha256": workbook_hash_after,
		"configured_job_count": configured_job_count,
		"latest_observed_jobs": latest_observed_jobs,
		"expected_workbook_jobs": expected_observed_jobs,
		"event_counts": dict(event_counts),
		"today_event_counts": dict(today_event_counts),
		"checks_passed": len(checks) - len(failed_checks),
		"checks_total": len(checks),
		"failed_checks": failed_checks,
		"checks": checks,
	}


def _select_run_for_verification(
	connection: sqlite3.Connection,
	*,
	profile: RadarProfile,
	output_scope: str | None,
) -> tuple[sqlite3.Row, dict[str, Any], bool]:
	"""Select a run without silently falling back past corrupt newer stats.

	Legacy stats omit ``output_scope`` and therefore describe a configured
	profile output.  For a scope-aware verification, a corrupt newer row is
	unsafe to skip: it could belong to the requested profile and scope.
	"""
	rows = connection.execute(
		"SELECT run_id, mode, status, started_at, completed_at, stats_json "
		"FROM radar_runs ORDER BY started_at DESC"
	).fetchall()
	if not rows:
		raise RadarAcceptanceError("岗位雷达数据库中没有运行记录")
	if output_scope is None:
		latest = rows[0]
		try:
			stats = json.loads(str(latest["stats_json"]))
		except (TypeError, ValueError):
			return latest, {}, False
		return latest, stats if isinstance(stats, dict) else {}, isinstance(stats, dict)

	for row in rows:
		try:
			stats = json.loads(str(row["stats_json"]))
		except (TypeError, ValueError) as exc:
			raise RadarAcceptanceError(
				"无法可靠按输出范围选择验收运行：较新的运行统计损坏 "
				f"({row['run_id']})"
			) from exc
		if not isinstance(stats, dict):
			raise RadarAcceptanceError(
				"无法可靠按输出范围选择验收运行：较新的运行统计不是对象 "
				f"({row['run_id']})"
			)
		run_fingerprint = stats.get("profile_fingerprint")
		if not isinstance(run_fingerprint, str) or not run_fingerprint.strip():
			raise RadarAcceptanceError(
				"无法可靠按输出范围选择验收运行：较新的运行缺少有效配置指纹 "
				f"({row['run_id']})"
			)
		if run_fingerprint != profile.fingerprint:
			continue
		actual_scope = stats.get("output_scope", "configured_profile")
		if not isinstance(actual_scope, str) or actual_scope not in {
			"configured_profile", "temporary_page_override",
		}:
			raise RadarAcceptanceError(
				"无法可靠按输出范围选择验收运行：匹配配置指纹的运行输出范围无效 "
				f"({row['run_id']})"
			)
		if actual_scope == output_scope:
			return row, stats, True
	raise RadarAcceptanceError(
		"找不到与当前配置指纹和输出范围一致的运行："
		f"profile={profile.name}, output_scope={output_scope}"
	)


def _natural_end_page_evidence(
	connection: sqlite3.Connection,
	*,
	run_ids: tuple[str, ...],
	expected_searches: set[str],
	require_visible_state: bool,
) -> dict[str, Any]:
	"""Assess terminal pagination from unfiltered page captures.

	The job-observation table is deliberately not used here: it omits irrelevant
	rows, so a short relevance-filtered page says nothing about the native BOSS
	page length.  Legacy runs without this table remain valid configured-window
	runs, but they cannot be certified as a natural platform end.
	"""
	base: dict[str, Any] = {
		"table_present": _table_exists(connection, "radar_page_captures"),
		"visible_state_required": require_visible_state,
		"missing_searches": [],
		"sequence_failures": [],
		"missing_visible_state": [],
		"missing_terminal_page_size": [],
		"terminal_full_searches": [],
		"page_rollbacks": [],
		"searches": [],
		"proven": False,
	}
	if not base["table_present"] or not run_ids:
		base["missing_searches"] = sorted(expected_searches)
		return base

	marks = ",".join("?" for _ in run_ids)
	rows = connection.execute(
		f"SELECT p.* FROM radar_page_captures p WHERE p.run_id IN ({marks})", run_ids,
	).fetchall()
	positions = {run_id: position for position, run_id in enumerate(run_ids)}
	rows.sort(key=lambda row: (positions[row["run_id"]], str(row["shard"]), str(row["query"]), int(row["page"])))
	by_search: dict[str, list[sqlite3.Row]] = {}
	for row in rows:
		key = f"{str(row['shard'])} / {str(row['query'])}"
		pages = by_search.setdefault(key, [])
		if pages and int(row["page"]) <= int(pages[-1]["page"]):
			# A later segment restarted a search. Do not fill its unseen tail with
			# higher pages left over from the earlier browser list.
			base["page_rollbacks"].append({"search": key, "restart_batch": int(row["page"])})
			pages.clear()
		pages.append(row)

	for key in sorted(expected_searches):
		pages = by_search.get(key, [])
		base["searches"].append({
			"search": key,
			"recorded_batches": len(pages),
			"raw_rows": sum(int(row["raw_job_count"]) for row in pages) if pages else None,
			"batches": [
				{"batch": int(row["page"]), "raw_rows": int(row["raw_job_count"]),
				 "page_size": row["page_size"], "has_more": bool(row["has_more"]),
				 "same_page_state": bool(row["same_page_state"])}
				for row in pages
			],
		})
		if not pages:
			base["missing_searches"].append(key)
			continue
		page_numbers = [int(row["page"]) for row in pages]
		terminal = pages[-1]
		sequence_ok = (
			page_numbers == list(range(1, page_numbers[-1] + 1))
			and int(terminal["has_more"]) == 0
			and all(int(row["has_more"]) == 1 for row in pages[:-1])
		)
		if not sequence_ok:
			base["sequence_failures"].append({
				"search": key,
				"pages": page_numbers,
				"terminal_has_more": bool(terminal["has_more"]),
			})
			continue
		if require_visible_state and any(int(row["same_page_state"]) != 1 for row in pages):
			base["missing_visible_state"].append(key)
		terminal_page_size = terminal["page_size"]
		if terminal_page_size is None:
			base["missing_terminal_page_size"].append(key)
		elif terminal_page_size is not None and int(terminal["raw_job_count"]) >= int(terminal_page_size):
			base["terminal_full_searches"].append({
				"search": key,
				"page": int(terminal["page"]),
				"raw_job_count": int(terminal["raw_job_count"]),
				"page_size": int(terminal_page_size),
			})

	base["proven"] = not any(
		base[name]
		for name in (
			"missing_searches",
			"sequence_failures",
			"missing_visible_state",
			"missing_terminal_page_size",
			"terminal_full_searches",
			"page_rollbacks",
		)
	)
	return base


def _read_only_connection(database_path: Path) -> sqlite3.Connection:
	connection = sqlite3.connect(f"{database_path.as_uri()}?mode=ro", uri=True)
	connection.row_factory = sqlite3.Row
	connection.execute("PRAGMA query_only = ON")
	return connection


def _scalar_count(connection: sqlite3.Connection, query: str, parameters: tuple[Any, ...] = ()) -> int:
	row = connection.execute(query, parameters).fetchone()
	return int(row[0]) if row is not None else 0


def _table_exists(connection: sqlite3.Connection, table_name: str) -> bool:
	row = connection.execute(
		"SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?",
		(table_name,),
	).fetchone()
	return row is not None


def _event_counts(connection: sqlite3.Connection, run_id: str) -> Counter[str]:
	return Counter(
		{
			str(row[0]): int(row[1])
			for row in connection.execute(
				"SELECT event_type, COUNT(*) FROM radar_events WHERE run_id = ? GROUP BY event_type",
				(run_id,),
			)
		}
	)


def _event_counts_for_runs(connection: sqlite3.Connection, run_ids: tuple[str, ...]) -> Counter[str]:
	if not run_ids:
		return Counter()
	marks = ",".join("?" for _ in run_ids)
	return Counter(
		{
			str(row[0]): int(row[1])
			for row in connection.execute(
				f"SELECT event_type, COUNT(*) FROM radar_events WHERE run_id IN ({marks}) GROUP BY event_type",
				run_ids,
			)
		}
	)


def _today_event_counts(
	connection: sqlite3.Connection,
	*,
	profile_fingerprint: str,
	start: datetime,
	end: datetime,
) -> Counter[str]:
	counts: Counter[str] = Counter()
	rows = connection.execute(
		"SELECT e.event_type, r.stats_json FROM radar_events AS e "
		"JOIN radar_runs AS r ON r.run_id = e.run_id "
		"WHERE e.event_type IN ('new', 'reposted') AND e.created_at >= ? AND e.created_at <= ?",
		(start.isoformat(), end.isoformat()),
	)
	for event_type, stats_json in rows:
		try:
			stats = json.loads(str(stats_json))
		except (TypeError, ValueError):
			continue
		if isinstance(stats, dict) and stats.get("profile_fingerprint") == profile_fingerprint:
			counts[str(event_type)] += 1
	return counts


def _configured_job_count(connection: sqlite3.Connection, profile: RadarProfile) -> int:
	query_marks = ",".join("?" for _ in profile.queries)
	city_marks = ",".join("?" for _ in profile.cities)
	return _scalar_count(
		connection,
		f"SELECT COUNT(DISTINCT j.id) FROM radar_jobs AS j "
		f"JOIN radar_job_sources AS s ON s.job_pk = j.id "
		f"WHERE s.query IN ({query_marks}) AND s.shard IN ({city_marks})",
		(*profile.queries, *(city.name for city in profile.cities)),
	)


def _inspect_workbook(workbook_path: Path, first_seen_by_job_id: dict[str, str]) -> dict[str, Any]:
	workbook = load_workbook(workbook_path, read_only=False, data_only=False, keep_links=False)
	try:
		sheet_names = list(workbook.sheetnames)
		header_mismatches: dict[str, list[Any]] = {}
		today_links: list[str] = []
		observed_links: list[str] = []
		summary: dict[str, Any] = {}
		first_seen_mismatches: list[dict[str, str]] = []
		job_sheet_name = "本轮首次发现" if "本轮首次发现" in sheet_names else "今日新增"
		if job_sheet_name in sheet_names:
			today_sheet = workbook[job_sheet_name]
			today_headers = [today_sheet.cell(1, column).value for column in range(1, 11)]
			if today_headers != _EXPECTED_JOB_HEADERS:
				header_mismatches[job_sheet_name] = today_headers
			today_links = _hyperlink_targets(today_sheet)
		if "本次扫描" in sheet_names:
			observed_sheet = workbook["本次扫描"]
			observed_headers = [observed_sheet.cell(1, column).value for column in range(1, 11)]
			if observed_headers != _EXPECTED_JOB_HEADERS:
				header_mismatches["本次扫描"] = observed_headers
			observed_links = _hyperlink_targets(observed_sheet)
			first_seen_mismatches = _first_seen_mismatches(observed_sheet, first_seen_by_job_id)
		if "扫描摘要" in sheet_names:
			summary_sheet = workbook["扫描摘要"]
			summary = {
				str(summary_sheet.cell(row, 1).value): summary_sheet.cell(row, 2).value
				for row in range(2, summary_sheet.max_row + 1)
				if summary_sheet.cell(row, 1).value not in (None, "搜索词", "未完成搜索")
			}
		formula_cells = [
			f"{sheet.title}!{cell.coordinate}"
			for sheet in workbook.worksheets
			for row in sheet.iter_rows()
			for cell in row
			if cell.data_type == "f"
		]
		return {
			"sheet_names": sheet_names,
			"header_mismatches": header_mismatches,
			"summary": summary,
			"today_links": today_links,
			"observed_links": observed_links,
			"first_seen_mismatches": first_seen_mismatches,
			"formula_cells": formula_cells,
		}
	finally:
		workbook.close()


def _hyperlink_targets(sheet: Any) -> list[str]:
	targets: list[str] = []
	for row in sheet.iter_rows():
		for cell in row:
			target = getattr(cell.hyperlink, "target", None)
			if target:
				targets.append(str(target))
	return targets


def _first_seen_mismatches(sheet: Any, first_seen_by_job_id: dict[str, str]) -> list[dict[str, str]]:
	mismatches: list[dict[str, str]] = []
	for row_number in range(2, sheet.max_row + 1):
		link_cell = sheet.cell(row_number, 10)
		target = getattr(link_cell.hyperlink, "target", None)
		if not target:
			continue
		job_id = _job_id_from_link(str(target))
		expected = first_seen_by_job_id.get(job_id)
		actual = str(sheet.cell(row_number, 2).value or "")
		if expected != actual:
			mismatches.append(
				{
					"job_id": job_id,
					"expected": str(expected or "missing"),
					"actual": actual,
				}
			)
	return mismatches


def _job_id_from_link(target: str) -> str:
	path = unquote(urlparse(target).path)
	return path.removeprefix("/job_detail/").removesuffix(".html")


def _is_job_link(target: str) -> bool:
	parsed = urlparse(target)
	host = (parsed.hostname or "").lower()
	return (
		parsed.scheme == "https"
		and (host == "zhipin.com" or host.endswith(".zhipin.com"))
		and parsed.path.startswith("/job_detail/")
		and parsed.path.endswith(".html")
		and bool(_job_id_from_link(target))
	)


def _parse_run_time(value: Any) -> datetime:
	try:
		parsed = datetime.fromisoformat(str(value))
	except (TypeError, ValueError) as exc:
		raise RadarAcceptanceError(f"运行时间格式无效: {value!r}") from exc
	if parsed.tzinfo is None:
		raise RadarAcceptanceError(f"运行时间缺少时区: {value!r}")
	return parsed


def _format_first_seen(value: str) -> str:
	return _parse_run_time(value).astimezone(_CHINA_TZ).strftime("%Y-%m-%d %H:%M:%S")


def _as_int(value: Any) -> int:
	try:
		return int(value)
	except (TypeError, ValueError):
		return -1


def _sha256(path: Path) -> str:
	digest = hashlib.sha256()
	with path.open("rb") as stream:
		for chunk in iter(lambda: stream.read(1024 * 1024), b""):
			digest.update(chunk)
	return digest.hexdigest()
