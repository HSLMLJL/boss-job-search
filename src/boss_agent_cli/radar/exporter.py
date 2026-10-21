from __future__ import annotations

import shutil
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from boss_agent_cli.radar.models import RadarEvent, ai_relevance_evidence
from boss_agent_cli.radar.scanner import RadarScanResult


_CHINA_TZ = timezone(timedelta(hours=8))
_HEADER_FILL = PatternFill("solid", fgColor="167D8D")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
_WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
_STATUS_LABELS = {
	"new": "新岗位",
	"reposted": "疑似重发",
	"observed": "本次扫描相关岗位（不代表新增）",
}
_FORMULA_PREFIXES = ("=", "+", "-", "@")


class RadarExportError(RuntimeError):
	"""Raised when the daily workbook cannot be written."""

	def __init__(self, message: str, *, saved_workbook: Path | None = None) -> None:
		super().__init__(message)
		self.saved_workbook = saved_workbook


@dataclass(frozen=True)
class RadarExportPaths:
	latest: Path
	last_successful: Path
	archive: Path
	run_archive: Path

	def to_dict(self) -> dict[str, str]:
		return {
			"latest_xlsx": str(self.latest.resolve()),
			"last_successful_xlsx": str(self.last_successful.resolve()),
			"archive_xlsx": str(self.archive.resolve()),
			"run_archive_xlsx": str(self.run_archive.resolve()),
		}


def export_radar_workbook(
	result: RadarScanResult,
	output_dir: Path,
	*,
	now: datetime | None = None,
) -> RadarExportPaths:
	local_now = (now or datetime.now(_CHINA_TZ)).astimezone(_CHINA_TZ)
	output_dir.mkdir(parents=True, exist_ok=True)
	archive_dir = output_dir / "archive"
	archive_dir.mkdir(parents=True, exist_ok=True)
	archive_path = archive_dir / f"{local_now:%Y-%m-%d}.xlsx"
	run_archive_dir = archive_dir / "runs" / f"{local_now:%Y-%m-%d}"
	run_archive_dir.mkdir(parents=True, exist_ok=True)
	existing_run_archives = sorted(run_archive_dir.glob(f"*-{result.run_id}.xlsx"))
	run_archive_path = (
		existing_run_archives[0]
		if existing_run_archives
		else run_archive_dir / f"{local_now:%H%M%S}-{result.run_id}.xlsx"
	)
	latest_path = output_dir / "latest.xlsx"
	last_successful_path = output_dir / "last-successful.xlsx"
	temporary_path = output_dir / f".{result.run_id}.xlsx.tmp"
	archive_temporary_path = output_dir / f".{result.run_id}.archive.xlsx.tmp"
	latest_temporary_path = output_dir / f".{result.run_id}.latest.xlsx.tmp"
	last_successful_temporary_path = output_dir / f".{result.run_id}.last-successful.xlsx.tmp"

	workbook = _build_workbook(result, local_now)
	try:
		workbook.save(temporary_path)
		temporary_path.replace(run_archive_path)
		shutil.copyfile(run_archive_path, archive_temporary_path)
		archive_temporary_path.replace(archive_path)
		shutil.copyfile(run_archive_path, latest_temporary_path)
		latest_temporary_path.replace(latest_path)
		# Keep one stable, user-facing workbook that failures and partial scans can
		# never replace. ``latest.xlsx`` remains the most recent audit attempt.
		if result.complete:
			shutil.copyfile(run_archive_path, last_successful_temporary_path)
			last_successful_temporary_path.replace(last_successful_path)
	except OSError as exc:
		# The immutable per-run archive is written before the replaceable daily
		# aliases. On Windows, an open latest.xlsx can prevent its replacement;
		# surface the completed archive so recovery never requires another scan.
		saved_workbook = run_archive_path if run_archive_path.is_file() else None
		raise RadarExportError(
			f"无法更新岗位雷达 Excel：{exc}",
			saved_workbook=saved_workbook,
		) from exc
	finally:
		workbook.close()
		if temporary_path.exists():
			temporary_path.unlink()
		if archive_temporary_path.exists():
			archive_temporary_path.unlink()
		if latest_temporary_path.exists():
			latest_temporary_path.unlink()
		if last_successful_temporary_path.exists():
			last_successful_temporary_path.unlink()
	return RadarExportPaths(
		latest=latest_path,
		last_successful=last_successful_path,
		archive=archive_path,
		run_archive=run_archive_path,
	)


def _build_workbook(result: RadarScanResult, local_now: datetime) -> Workbook:
	workbook = Workbook()
	jobs_sheet = workbook.active
	jobs_sheet.title = "本轮首次发现" if result.event_window else "今日新增"
	_write_jobs_sheet(jobs_sheet, result)
	observed_sheet = workbook.create_sheet("本次扫描")
	_write_observed_jobs_sheet(observed_sheet, result)
	summary_sheet = workbook.create_sheet("扫描摘要")
	_write_summary_sheet(summary_sheet, result, local_now)
	workbook.calculation.fullCalcOnLoad = True
	return workbook


def _write_jobs_sheet(sheet: Any, result: RadarScanResult) -> None:
	headers = _job_headers()
	sheet.append(headers)
	_style_header(sheet, 1, len(headers))
	display_events = result.reportable_new_events
	if result.mode == "baseline":
		# A baseline defines the comparison set; it must never populate a sheet
		# named "今日新增" with historical inventory. The same jobs remain fully
		# available, with clickable links, on "本次扫描".
		display_events = []
	for event in sorted(display_events, key=_event_sort_key):
		_append_job_row(sheet, event)

	if not display_events:
		sheet.append([result._note()])
		sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
		sheet.cell(2, 1).alignment = Alignment(wrap_text=True)
	_finish_jobs_sheet(sheet)


def _write_observed_jobs_sheet(sheet: Any, result: RadarScanResult) -> None:
	headers = _job_headers()
	sheet.append(headers)
	_style_header(sheet, 1, len(headers))
	display_events = [RadarEvent(event_type="observed", job=job) for job in result.observed_jobs]
	for event in sorted(display_events, key=_event_sort_key):
		_append_job_row(sheet, event)
	if not display_events:
		sheet.append(["本次扫描未发现相关岗位；若覆盖不完整，请查看“扫描摘要”。"])
		sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
		sheet.cell(2, 1).alignment = Alignment(wrap_text=True)
	_finish_jobs_sheet(sheet)


def _job_headers() -> list[str]:
	return [
		"状态",
		"首次发现时间",
		"岗位",
		"公司",
		"城市/区域",
		"薪资",
		"经验",
		"学历",
		"AI相关证据",
		"BOSS链接",
	]


def _append_job_row(sheet: Any, event: RadarEvent) -> None:
	job = event.job
	discovered = job.observed_at.astimezone(_CHINA_TZ).strftime("%Y-%m-%d %H:%M:%S")
	location = "·".join(value for value in (job.city, job.district) if value)
	sheet.append(
		[
			_STATUS_LABELS.get(event.event_type, event.event_type),
			discovered,
			_safe_excel_text(job.title),
			_safe_excel_text(job.company),
			_safe_excel_text(location),
			_safe_excel_text(job.salary),
			_safe_excel_text(job.experience),
			_safe_excel_text(job.education),
			_safe_excel_text("、".join(ai_relevance_evidence(job))),
			"打开岗位" if job.job_url else "链接不可用",
		]
	)
	link_cell = sheet.cell(row=sheet.max_row, column=10)
	if job.job_url:
		link_cell.hyperlink = job.job_url
		link_cell.style = "Hyperlink"


def _finish_jobs_sheet(sheet: Any) -> None:

	sheet.freeze_panes = "A2"
	sheet.auto_filter.ref = f"A1:J{max(sheet.max_row, 1)}"
	widths = [30, 21, 32, 25, 22, 14, 12, 10, 28, 14]
	for index, width in enumerate(widths, start=1):
		sheet.column_dimensions[get_column_letter(index)].width = width
	for row in sheet.iter_rows(min_row=2):
		for cell in row:
			cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_summary_sheet(sheet: Any, result: RadarScanResult, local_now: datetime) -> None:
	coverage = result.to_dict()["coverage"]
	totals = result.scan_totals or {
		"requests": result.requests,
		"elapsed_seconds": result.elapsed_seconds,
		"pages": result.pages,
		"jobs_seen": result.jobs_seen,
		"relevant_seen": result.relevant_seen,
		"unique_seen": result.unique_seen,
		"query_metrics": result.query_metrics,
		"event_counts": {
			"new": sum(event.event_type == "new" for event in result.new_events),
			"reposted": sum(event.event_type == "reposted" for event in result.new_events),
		},
	}
	event_counts = totals.get("event_counts", {})
	if not isinstance(event_counts, dict):
		event_counts = {}
	rows = [
		("配置", result.profile.name),
		("运行ID", result.run_id),
		("运行模式", result.mode_label),
		*( [("本轮窗口", f"{result.event_window['start_at']} 至 {result.event_window['end_at']}（非自然日‘今日新增’）")] if result.event_window else [] ),
		("生成时间", local_now.strftime("%Y-%m-%d %H:%M:%S UTC+8")),
		("覆盖状态", result.coverage_label),
		("配置城市", len(result.profile.cities)),
		("已完成搜索城市", len(result.successful_cities)),
		("覆盖待核城市", "、".join(result.coverage_review_cities) or "无"),
		("失败/未完成城市", "、".join(result.failed_cities) or "无"),
		("搜索组合", f"{coverage['completed_searches']} / {coverage['planned_searches']}"),
		("继承上次已完成搜索", coverage["resumed_completed_searches"]),
		("请求数", totals["requests"]),
		("扫描耗时", f"{float(totals['elapsed_seconds']) / 60:.1f} 分钟"),
		("完成页数", totals["pages"]),
		("达到分页上限的搜索", len(result.page_cap_reached_searches)),
		("本次读取岗位", totals["jobs_seen"]),
		("本次相关岗位", totals["relevant_seen"]),
		("本次唯一相关岗位", totals["unique_seen"]),
		("当前清单岗位", len(result.observed_jobs)),
		("本次新增岗位", int(event_counts.get("new", 0))),
		("本次疑似重发", int(event_counts.get("reposted", 0))),
		(("本轮累计新增岗位" if result.event_window else "今日累计新增岗位"), sum(event.event_type == "new" for event in result.reportable_new_events)),
		(("本轮累计疑似重发" if result.event_window else "今日累计疑似重发"), sum(event.event_type == "reposted" for event in result.reportable_new_events)),
		("停止代码", result.stop_code or "无"),
		("停止原因", result.stop_reason or "无"),
		("结论", result._note()),
		("城市分页计划", _format_city_page_plan(result)),
		("计划页面请求", coverage["planned_page_requests"]),
	]
	sheet.append(["项目", "结果"])
	_style_header(sheet, 1, 2)
	for key, value in rows:
		sheet.append([key, _safe_excel_value(value)])
	if not result.complete:
		for cell in sheet[sheet.max_row]:
			cell.fill = _WARNING_FILL
			cell.font = Font(bold=True)

	start_row = sheet.max_row + 3
	sheet.cell(start_row, 1, "搜索词")
	sheet.cell(start_row, 2, "读取岗位")
	sheet.cell(start_row, 3, "相关岗位")
	sheet.cell(start_row, 4, "独有贡献")
	sheet.cell(start_row, 5, "本配置首词未覆盖")
	_style_header(sheet, start_row, 5)
	total_query_metrics = totals.get("query_metrics", {})
	if not isinstance(total_query_metrics, dict):
		total_query_metrics = {}
	for query in result.profile.queries:
		metrics = total_query_metrics.get(query, {})
		if not isinstance(metrics, dict):
			metrics = {}
		sheet.append(
			[
				_safe_excel_text(query),
				metrics.get("jobs_seen", 0),
				metrics.get("relevant_seen", 0),
				metrics.get("unique_contributed", 0),
				metrics.get("not_in_primary_query", 0),
			]
		)

	if result.truncated_searches or result.failed_searches:
		problem_row = sheet.max_row + 3
		sheet.cell(problem_row, 1, "未完成搜索")
		sheet.cell(problem_row, 2, "原因")
		_style_header(sheet, problem_row, 2)
		for item in result.failed_searches:
			sheet.append(
				[
					_safe_excel_text(item.split(":", 1)[0]),
					_safe_excel_text(item.split(":", 1)[-1].strip()),
				]
			)
		failed_keys = {item.split(":", 1)[0] for item in result.failed_searches}
		for item in result.truncated_searches:
			if item not in failed_keys:
				sheet.append([_safe_excel_text(item), "请求上限、分页上限或安全停止后未继续"])

	# Explicit white avoids transparent blank cells rendering as black in some
	# viewers while preserving semantic header and warning fills.
	for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=5):
		for cell in row:
			if cell.fill.fill_type is None:
				cell.fill = _WHITE_FILL
	sheet.freeze_panes = "A2"
	sheet.column_dimensions["A"].width = 32
	sheet.column_dimensions["B"].width = 70
	sheet.column_dimensions["C"].width = 16
	sheet.column_dimensions["D"].width = 16
	sheet.column_dimensions["E"].width = 16
	for row in sheet.iter_rows():
		for cell in row:
			# The first summary table uses a deliberately wide result column for
			# explanatory text.  Left-aligning its values keeps small numeric
			# results visually attached to their labels instead of appearing in a
			# phantom third column at the far edge of that wide cell.
			horizontal = "left" if cell.column == 2 and cell.row > 1 and cell.row < start_row else None
			cell.alignment = Alignment(horizontal=horizontal, vertical="top", wrap_text=True)


def _event_sort_key(event: RadarEvent) -> tuple[int, float, str, str]:
	status_order = 0 if event.event_type == "new" else 1
	return (
		status_order,
		-event.job.observed_at.timestamp(),
		event.job.city,
		event.job.title,
	)


def _format_city_page_plan(result: RadarScanResult) -> str:
	if result.profile.query_page_limits:
		search_limits = result.search_page_limits or result.profile.search_page_limits
		groups: dict[int, list[str]] = {}
		for search, pages in search_limits.items():
			groups.setdefault(pages, []).append(search)
		return "；".join(
			f"{pages}页：{'、'.join(searches)}"
			for pages, searches in sorted(groups.items(), reverse=True)
		)
	limits = result.page_limits or result.profile.city_page_limits
	grouped: dict[int, list[str]] = {}
	for city in result.profile.cities:
		grouped.setdefault(limits[city.name], []).append(city.name)
	return "；".join(
		f"{pages}页：{'、'.join(cities)}"
		for pages, cities in sorted(grouped.items(), key=lambda item: item[0], reverse=True)
	)


def _style_header(sheet: Any, row: int, columns: int) -> None:
	for column in range(1, columns + 1):
		cell = sheet.cell(row=row, column=column)
		cell.fill = _HEADER_FILL
		cell.font = _HEADER_FONT
		cell.alignment = Alignment(horizontal="center", vertical="center")


def _safe_excel_text(value: Any) -> str:
	text = str(value or "")
	if text.lstrip().startswith(_FORMULA_PREFIXES):
		return f"'{text}"
	return text


def _safe_excel_value(value: Any) -> Any:
	return _safe_excel_text(value) if isinstance(value, str) else value
