import asyncio
import csv
import json
from copy import deepcopy

from aiohttp.test_utils import TestClient, TestServer
from click.testing import CliRunner
from openpyxl import load_workbook
import pytest

from job_radar.cli import cli
from job_radar.config import DEFAULT_CONFIG, TaskConfig
from job_radar.exporter import export_run, public_job
from job_radar.filters import evaluate_job
from job_radar.server import create_app
from job_radar.service import Workspace, collection_lock, execute_run
from job_radar.store import JobStore


def config(**changes):
	raw = deepcopy(DEFAULT_CONFIG)
	raw.update(changes)
	return TaskConfig.from_dict(raw)


def job(**changes):
	return {
		"job_key": "test_1",
		"title": "产品经理",
		"company": "示例公司",
		"city": "上海",
		"salary": "15-25K",
		"queries": ["产品经理"],
		"matched": True,
		"detail_status": "not_requested",
		"description": "",
		**changes,
	}


@pytest.mark.parametrize(
	"change",
	[
		{"keywords": []},
		{"cities": ["不存在城市"]},
		{"details": "invalid"},
		{"limits": {"requests": True}},
		{"export": {"fields": ["_security_id"]}},
		{"filters": {"experience": "invalid"}},
		{"filters": {"unknown": "guess"}},
	],
)
def test_invalid_configuration_is_rejected(change):
	with pytest.raises(ValueError):
		TaskConfig.from_dict(change)


def test_configuration_is_independent_and_roundtrips():
	parsed = TaskConfig.from_dict({"keywords": ["运营", " 运营 ", "数据分析"]})
	assert parsed.keywords == ["运营", "数据分析"]
	assert TaskConfig.from_dict(parsed.to_dict()) == parsed
	parsed.filters["exclude"].append("销售")
	assert DEFAULT_CONFIG["filters"]["exclude"] == []


def test_unknown_jd_is_distinct_from_negative_evidence():
	filters = {**DEFAULT_CONFIG["filters"], "include": ["增长"], "unknown": "exclude"}
	assert evaluate_job(job(), filters, detail_pending=True)[0]
	assert not evaluate_job(job(), filters)[0]
	assert evaluate_job(job(description="负责用户增长", detail_status="fetched"), filters)[0]
	assert not evaluate_job(job(description="负责客服", detail_status="fetched"), filters)[0]
	filters.update(include=[], salary="10-20K")
	assert evaluate_job(job(), filters)[0]
	assert not evaluate_job(job(salary="30-40K"), filters)[0]


def test_history_tracks_real_listing_changes_and_preserves_notes(tmp_path):
	store = JobStore(tmp_path)
	first = store.create_run(config().to_dict())
	saved = store.upsert_job(first["id"], job())
	store.update_job_status("test_1", "已投", "周一跟进")
	store.upsert_job(first["id"], job(description="完整 JD", skills=["分析"], queries=["运营"]))
	assert store.get_job(first["id"], "test_1")["queries"] == ["产品经理", "运营"]
	second = store.create_run(config().to_dict())
	again = store.upsert_job(second["id"], job())
	assert again["first_seen"] == saved["first_seen"]
	assert (again["change"], again["status"], again["notes"]) == ("seen", "已投", "周一跟进")
	third = store.create_run(config().to_dict())
	assert store.upsert_job(third["id"], job(salary="20-30K"))["change"] == "changed"


def test_exports_preserve_complete_jd_and_are_safe_and_immutable(tmp_path):
	store = JobStore(tmp_path)
	cfg = config(
		export={"formats": ["xlsx", "csv", "json"], "fields": ["title", "description", "url", "notes"], "scope": "all"}
	)
	run = store.create_run(cfg.to_dict(), demo=True)
	long_jd = "完整岗位说明\n" * 6000
	store.upsert_job(
		run["id"],
		job(
			title="=1+1",
			description=long_jd,
			url="https://www.zhipin.com/job_detail/example.html?securityId=private",
			_security_id="private",
		),
	)
	store.update_job_status("test_1", "考虑", '=HYPERLINK("https://example.com")')
	store.update_run(run["id"], status="partial", message="仍有未访问页面")
	files = export_run(store, run["id"])
	paths = {
		item["name"].split(".")[-1]: store.root / "exports" / item["relative"]
		for item in files
		if item["name"] != "运行说明.json"
	}
	data = json.loads(paths["json"].read_text(encoding="utf-8"))
	assert data["metadata"]["status"] == "partial" and data["metadata"]["demo"]
	assert data["jobs"][0]["description"] == long_jd
	assert "private" not in paths["json"].read_text(encoding="utf-8")
	with paths["csv"].open(encoding="utf-8-sig", newline="") as stream:
		rows = list(csv.reader(stream))
	assert rows[1][0] == "'=1+1" and rows[1][1] == long_jd
	book = load_workbook(paths["xlsx"])
	assert book["岗位清单"]["A2"].data_type == "s"
	assert "".join(row[3] for row in book["长文本"].iter_rows(min_row=2, values_only=True)) == long_jd
	second = export_run(store, run["id"])
	assert files[0]["relative"] != second[0]["relative"]
	assert all(path.exists() for path in paths.values())
	assert public_job(job(url="javascript:alert(1)", _security_id="secret"))["url"] == ""


def test_execute_marks_export_pending_only_while_writing_files(tmp_path, monkeypatch):
	store = JobStore(tmp_path)
	run = store.create_run(config().to_dict(), demo=True)
	phases = []

	def fake_run_task(_config, current_store, _source, *, run_id, stop_event):
		phases.append(("collect", current_store.get_run(run_id).get("export_pending", False)))
		current_store.update_run(run_id, status="completed")

	def fake_export(current_store, run_id):
		phases.append(("export", current_store.get_run(run_id).get("export_pending", False)))
		return []

	monkeypatch.setattr("job_radar.runner.run_task", fake_run_task)
	monkeypatch.setattr("job_radar.service.export_run", fake_export)

	execute_run(store, run["id"])

	assert phases == [("collect", False), ("export", True)]
	assert not store.get_run(run["id"])["export_pending"]


def test_ui_starts_when_bridge_daemon_cannot_start(tmp_path, monkeypatch):
	def unavailable():
		raise RuntimeError("daemon unavailable")

	monkeypatch.setattr("job_radar.cli._start_bridge", unavailable)
	monkeypatch.setattr("aiohttp.web.run_app", lambda *_args, **_kwargs: None)

	result = CliRunner().invoke(cli, ["--data-dir", str(tmp_path), "ui", "--no-open"])

	assert result.exit_code == 0, result.output
	assert "合成示例和本地记录仍可使用" in result.output


def test_source_construction_failure_preserves_export(tmp_path, monkeypatch):
	store = JobStore(tmp_path)
	run = store.create_run(config().to_dict())

	def fail(*args):
		raise RuntimeError("private token")

	monkeypatch.setattr("job_radar.source.make_source", fail)
	with pytest.raises(RuntimeError):
		execute_run(store, run["id"])
	result = store.get_run(run["id"])
	assert result["status"] == "failed" and result["exports"]
	assert "private token" not in result["message"]


def test_cli_demo_is_offline_and_delivers_all_formats(tmp_path, monkeypatch):
	def forbidden(*args):
		raise AssertionError("demo must never construct real platform source")

	monkeypatch.setattr("job_radar.source.make_source", forbidden)
	result = CliRunner().invoke(cli, ["--data-dir", str(tmp_path), "demo"])
	assert result.exit_code == 0, result.output
	run = JobStore(tmp_path).list_runs()[0]
	assert run["demo"] and run["status"] == "completed"
	assert len(run["exports"]) == 4


def test_workspace_recovers_orphans_but_respects_active_process_lock(tmp_path):
	store = JobStore(tmp_path)
	run = store.create_run(config().to_dict())
	store.update_run(run["id"], status="running", export_pending=True)
	with collection_lock(store.root):
		Workspace(tmp_path)
		assert store.get_run(run["id"])["status"] == "running"
	Workspace(tmp_path)
	recovered = store.get_run(run["id"])
	assert recovered["status"] == "stopped" and not recovered["export_pending"]


def test_server_contract_and_local_boundary(tmp_path, monkeypatch):
	# Browser connectivity is global to this machine, not scoped to data_dir.
	monkeypatch.setattr("job_radar.server.auth_status", lambda _: {"ready": False, "message": "未连接"})

	async def exercise():
		client = TestClient(TestServer(create_app(tmp_path)))
		await client.start_server()
		try:
			response = await client.get("/api/bootstrap")
			body = await response.json()
			token = {"X-Job-Radar-Token": body["csrf_token"]}
			assert response.status == 200 and not body["auth"]["ready"]
			assert (await client.post("/api/templates", json={"config": {}})).status == 403
			assert (await client.get("/api/bootstrap", headers={"Host": "evil.test"})).status == 403
			assert (await client.get("/api/bootstrap", headers={"Origin": "https://evil.test"})).status == 403
			saved = await client.post("/api/templates", json={"config": {"keywords": ["运营"]}}, headers=token)
			assert saved.status == 200
			response = await client.post("/api/runs", json={"config": {}, "demo": True}, headers=token)
			assert response.status == 201
			run_id = (await response.json())["id"]
			for _ in range(100):
				detail = await (await client.get(f"/api/runs/{run_id}")).json()
				if detail["exports"]:
					break
				await asyncio.sleep(0.03)
			assert detail["run"]["status"] == "completed" and detail["exports"]
			assert "checkpoint" not in detail["run"]
			assert detail["jobs"] and all("_security_id" not in row for row in detail["jobs"])
			listed = await (await client.get("/api/runs")).json()
			assert "checkpoint" not in listed["runs"][0]
			response = await client.get(detail["exports"][0]["url"])
			assert response.status == 200 and "attachment" in response.headers["Content-Disposition"]
			assert (await client.get("/static/missing.js")).status == 404
		finally:
			await client.close()

	asyncio.run(exercise())
